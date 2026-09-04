from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "RCA-2026-001_Validador-sin-conexion_Fundo-Santa-Isabel.docx"

RCA_ID = "RCA-2026-001"
REPORT_TITLE = 'Dos eventos "Validador sin conexión"'
LOCATION = "Fundo Santa Isabel"

BLACK = "000000"
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "2F6B4F"
CAUTION = "7A5A00"
RISK = "9B1C1C"
LINE = "D7DDE5"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: str = BLACK,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, attrs in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key, value in attrs.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def apply_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Las columnas deben sumar {CONTENT_WIDTH_DXA} DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, **CELL_MARGINS)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, *, color: str = LINE, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def add_paragraph_bottom_border(paragraph, *, color: str = NAVY, size: int = 14, space: int = 7) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str, display: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    style_specs = {
        "Title": (23, BLACK, 0, 4),
        "Subtitle": (14, MUTED, 0, 16),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in style_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_numbering_definition(doc: Document, *, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def add_bullets(doc: Document, items: list[str], num_id: int) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        apply_numbering(p, num_id)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str], num_id: int) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        apply_numbering(p, num_id)
        p.add_run(item)


def add_label_paragraph(doc: Document, label: str, value: str, *, after: float = 2) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, bold=True)
    r = p.add_run(value)
    set_run_font(r, size=10.5)


def add_table_text(cell, text: str, *, bold: bool = False, color: str = BLACK, size: float = 9.5, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)


def add_callout(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(table, [CONTENT_WIDTH_DXA])
    remove_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_border(
        cell,
        left={"val": "single", "sz": "22", "color": BLUE, "space": "0"},
        top={"val": "single", "sz": "4", "color": LINE, "space": "0"},
        bottom={"val": "single", "sz": "4", "color": LINE, "space": "0"},
        right={"val": "single", "sz": "4", "color": LINE, "space": "0"},
    )
    first = cell.paragraphs[0]
    first.clear()
    for idx, (label, value, color) in enumerate(rows):
        p = first if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(5 if idx < len(rows) - 1 else 0)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=10.5, color=color, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_header_footer(section) -> None:
    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = hp.add_run(f"{RCA_ID}  |  Monitoreo IoT Petróleo")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = hp.add_run("\tFINAL")
    set_run_font(right, size=8.5, color=GREEN, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = fp.add_run(f"{LOCATION}  |  Emisión 23-08-2026")
    set_run_font(left, size=8.5, color=MUTED)
    right = fp.add_run("\tPágina ")
    set_run_font(right, size=8.5, color=MUTED)
    add_field(fp, "PAGE", "1")
    middle = fp.add_run(" de ")
    set_run_font(middle, size=8.5, color=MUTED)
    add_field(fp, "NUMPAGES", "1")


def add_timeline_table(doc: Document) -> None:
    headers = ["Evento", "Primera ocurrencia", "Segunda ocurrencia"]
    rows = [
        ("Auditoría validator_link_lost", "19:19:46.675", "20:12:41.826"),
        ("manual_mode_aborted / service_stopping; relé 1 → 0", "19:19:46.682", "20:12:41.835"),
        ("Sistema registra apagado ordenado", "19:19:48", "20:12:43"),
        ("Nuevo boot", "19:20:04", "22:05:53"),
        ("Edge vuelve a asignarse", "19:20:43.669", "22:06:31.441"),
        ("Modo manual se reanuda", "19:20:47.139", "22:06:34.927"),
        ("Alerta pendiente se sincroniza a la web", "19:20:49", "22:06:37"),
    ]
    table = doc.add_table(rows=1, cols=3)
    apply_table_geometry(table, [5040, 2160, 2160])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        add_table_text(cell, text, bold=True, color=NAVY, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for event, first, second in rows:
        cells = table.add_row().cells
        add_table_text(cells[0], event, size=9.5)
        add_table_text(cells[1], first, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_table_text(cells[2], second, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    apply_table_geometry(table, [5040, 2160, 2160])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_control_table(doc: Document) -> None:
    rows = [
        ("Código", RCA_ID, "Versión", "1.0"),
        ("Estado", "Final", "Fecha de emisión", "23-08-2026"),
        ("Fecha del incidente", "22-08-2026", "Ubicación", LOCATION),
        ("Serie", "RCA-AAAA-NNN", "Siguiente correlativo", "RCA-2026-002"),
    ]
    table = doc.add_table(rows=0, cols=4)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            label = idx % 2 == 0
            if label:
                set_cell_shading(cells[idx], LIGHT_GRAY)
            add_table_text(cells[idx], value, bold=label, color=NAVY if label else BLACK, size=9.5)
    apply_table_geometry(table, [1680, 3000, 1800, 2880])
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_status_table(doc: Document) -> None:
    rows = [
        ("fuel-edge", "Activo", "NRestarts=0 en el boot actual"),
        ("mosquitto", "Activo", "NRestarts=0 en el boot actual"),
        ("fuel-edge-web", "Activo", "Servicio operativo"),
        ("Validador", "Conectado", "10.42.0.227 → 10.42.0.1:8883"),
        ("GPIO24", "Alto", "Pull-up activo; sin orden de apagado"),
        ("Fuel Edge", "Sin warnings", "Journal del boot actual"),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ["Componente", "Estado", "Evidencia"]
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_GRAY)
        add_table_text(cell, value, bold=True, color=NAVY, size=9.5)
    set_repeat_table_header(table.rows[0])
    for component, status, evidence in rows:
        cells = table.add_row().cells
        add_table_text(cells[0], component, size=9.5)
        add_table_text(cells[1], status, bold=True, color=GREEN, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_table_text(cells[2], evidence, size=9.5)
    apply_table_geometry(table, [1980, 1800, 5580])
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_document() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False
    settings = doc.settings._element
    for marker in list(settings.findall(qn("w:evenAndOddHeaders"))):
        settings.remove(marker)
    configure_header_footer(section)
    bullet_id = add_numbering_definition(doc, ordered=False)
    five_whys_id = add_numbering_definition(doc, ordered=True)

    props = doc.core_properties
    props.title = f"{RCA_ID} - {REPORT_TITLE}"
    props.subject = "Análisis de causa raíz de dos eventos de desconexión del validador"
    props.author = "Monitoreo IoT Petróleo"
    props.keywords = "RCA, validador, MQTT, UPS, GPIO24, Fundo Santa Isabel"
    props.comments = "Documento final - correlativo RCA anual"

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(4)
    r = kicker.add_run("ANÁLISIS DE CAUSA RAÍZ")
    set_run_font(r, size=9.5, color=BLUE, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run(RCA_ID)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(REPORT_TITLE)
    location = doc.add_paragraph()
    location.paragraph_format.space_after = Pt(16)
    r = location.add_run(f"Sistema de monitoreo y control de combustible | {LOCATION}")
    set_run_font(r, size=11, color=DARK_BLUE, bold=True)

    add_label_paragraph(doc, "Fecha del incidente", "22 de agosto de 2026")
    add_label_paragraph(doc, "Fecha de emisión", "23 de agosto de 2026")
    add_label_paragraph(doc, "Estado", "Final")
    add_label_paragraph(doc, "Versión", "1.0")
    add_label_paragraph(doc, "Clasificación", "Operacional - energía y conectividad", after=8)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    add_paragraph_bottom_border(rule)

    doc.add_heading("Resumen ejecutivo", level=1)
    intro = doc.add_paragraph()
    intro.add_run("Las dos alertas del 22-08-2026 no prueban dos fallas autónomas del validador RFID. ").bold = True
    intro.add_run(
        "En ambos casos, el enlace MQTT cayó porque la Raspberry PLC/edge estaba entrando en una secuencia de apagado ordenado."
    )
    add_callout(
        doc,
        [
            ("Causa inmediata confirmada", "apagado gestionado del controlador edge.", GREEN),
            ("Disparador físico más probable", "señal UPS/GPIO24 por pérdida o interrupción de la alimentación principal.", CAUTION),
            ("Defecto contribuyente", "Fuel Edge 0.3.15 clasifica el cierre MQTT intencional como desconexión inesperada del validador.", RISK),
            ("Resultado de seguridad", "el modo manual fue abortado y el relé se desenergizó en menos de 10 ms.", GREEN),
        ],
    )

    doc.add_heading("Nivel de certeza", level=2)
    add_bullets(
        doc,
        [
            "Apagado del edge como causa inmediata de ambas alertas: confirmado.",
            "Generación de la alerta dentro del cierre normal del servicio: confirmado por auditoría y código desplegado.",
            "Señal UPS por pérdida de alimentación principal: confianza alta, pero no absoluta por falta de journal persistente de los boots afectados.",
        ],
        bullet_id,
    )

    doc.add_page_break()
    doc.add_heading("Control documental y correlativo", level=2)
    add_control_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.add_run(
        "La serie usa el formato RCA-AAAA-NNN, con secuencia anual de tres dígitos. Este documento inaugura la serie 2026; el próximo análisis será RCA-2026-002."
    )

    doc.add_heading("1. Alcance y fuentes revisadas", level=1)
    add_bullets(
        doc,
        [
            "Exportación operacional sanitizada base-datos-fundo-santa-isabel-2026-08-23.json.",
            "Auditoría durable y outbox de /var/lib/fuel-edge/edge.db.",
            "Registro de boots y apagados del sistema operativo mediante last -x -F.",
            "Código fuente local y versión Fuel Edge 0.3.15 desplegada en producción.",
            "Configuración de device-tree para GPIO23/GPIO24 y servicios rpishutdown.",
            "Estado actual de fuel-edge, mosquitto, fuel-edge-web, conexiones TCP y GPIO24.",
        ],
        bullet_id,
    )
    p = doc.add_paragraph()
    p.add_run("Método: ").bold = True
    p.add_run(
        "correlación temporal entre alertas, auditoría del dominio, secuencia service_stopping, boots del host, brechas de telemetría y comportamiento del transporte MQTT."
    )

    doc.add_heading("2. Cronología de los incidentes", level=1)
    p = doc.add_paragraph()
    p.add_run("Zona horaria: ").bold = True
    p.add_run("America/Santiago (UTC-4). Las marcas originales de auditoría están en UTC.")
    add_timeline_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Lectura operacional. ").bold = True
    p.add_run(
        "La primera alerta quedó en la outbox durante el reinicio y se sincronizó después de recuperar el servicio. La segunda permaneció retenida durante casi dos horas y también llegó a la web sólo después del siguiente boot."
    )

    doc.add_heading("3. Hallazgos técnicos", level=1)
    doc.add_heading("3.1 Ambas alertas ocurrieron dentro del apagado", level=2)
    p = doc.add_paragraph()
    p.add_run(
        "En las dos ocurrencias, validator_link_lost fue seguido entre 7 y 10 ms después por manual_mode_aborted con reason=service_stopping, transición manual_mode → locked y relay_energized=0. "
    )
    p.add_run("Esta secuencia demuestra que el incidente de enlace se registró durante la detención del proceso.").bold = True

    doc.add_heading("3.2 Los boots confirman dos apagados ordenados", level=2)
    add_bullets(
        doc,
        [
            "Primer ciclo: shutdown 19:19:48; nuevo boot 19:20:04. Tiempo entre marcas: 16 s.",
            "Segundo ciclo: shutdown 20:12:43; nuevo boot 22:05:53. Tiempo entre marcas: 1 h 53 min 10 s.",
        ],
        bullet_id,
    )
    p = doc.add_paragraph()
    p.add_run(
        "Un corte abrupto sin respaldo o un watchdog duro no producirían la misma combinación de service_stopping, cierre seguro del relé y marca ordenada de shutdown."
    )

    doc.add_heading("3.3 La telemetría acompaña los periodos sin edge", level=2)
    add_bullets(
        doc,
        [
            "Primera ventana: OCIO 19:19:11.074 → 19:20:59.669; brecha de 1 min 48,595 s y nivel estable en 1.045,7 L.",
            "Segunda ventana: OCIO 20:12:26.914 → 22:06:47.451; brecha de 1 h 54 min 20,537 s, coherente con el tiempo sin controlador.",
            "No existe un despacho de combustible coincidente con las dos alertas.",
        ],
        bullet_id,
    )

    doc.add_heading("4. Análisis de causa raíz", level=1)
    doc.add_heading("4.1 Causa inmediata", level=2)
    p = doc.add_paragraph()
    p.add_run("Confirmada: ").bold = True
    p.add_run(
        "la Raspberry PLC/edge entró en apagado ordenado y, como consecuencia, cerró su transporte MQTT y dejó fuera de servicio al broker local durante ambos intervalos."
    )

    doc.add_heading("4.2 Disparador físico más probable", level=2)
    p = doc.add_paragraph()
    p.add_run("Alta confianza: ").bold = True
    p.add_run(
        "el circuito UPS activó la entrada GPIO24 al detectar pérdida o interrupción de la alimentación principal. La configuración instalada es:"
    )
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    r = code.add_run("dtoverlay=gpio-poweroff,gpiopin=23,active_low\n")
    set_run_font(r, name="Courier New", size=9.5, color=NAVY)
    r = code.add_run("dtoverlay=gpio-shutdown,gpio_pin=24,gpio_pull=up")
    set_run_font(r, name="Courier New", size=9.5, color=NAVY)
    p = doc.add_paragraph()
    p.add_run(
        "El kernel expone soc:shutdown_button@18; 0x18 equivale a GPIO24. El overlay genera KEY_POWER y systemd-logind inicia el apagado seguro."
    )

    doc.add_heading("4.3 Defecto contribuyente de clasificación", level=2)
    p = doc.add_paragraph()
    p.add_run("Fuel Edge 0.3.15 ").bold = True
    p.add_run(
        "no distingue entre desconexión MQTT inesperada y cierre intencional durante poweroff, reboot o detención del servicio."
    )
    steps_id = add_numbering_definition(doc, ordered=True)
    add_numbered(
        doc,
        [
            "MqttValidatorTransport.close() marca _started=False.",
            "Publica estado MQTT offline y llama client.disconnect().",
            "Paho ejecuta _on_disconnect().",
            "_on_disconnect() invoca siempre el handler de falla, sin evaluar cierre intencional ni reason_code.",
            "El handler aplica validator_link_lost y crea la alerta de prioridad alta.",
        ],
        steps_id,
    )
    p = doc.add_paragraph()
    p.add_run("Efecto: ").bold = True
    p.add_run(
        "la alarma atribuye el síntoma al validador y oculta que el origen fue el apagado del controlador edge."
    )

    doc.add_heading("4.4 Cinco porqués", level=2)
    add_numbered(
        doc,
        [
            "¿Por qué apareció la alerta? Porque Fuel Edge aplicó validator_link_lost.",
            "¿Por qué aplicó ese evento? Porque el callback MQTT invoca siempre el handler de falla.",
            "¿Por qué se desconectó MQTT? Porque Fuel Edge se cerraba durante el apagado del host.",
            "¿Por qué se apagó el host? Con alta probabilidad, GPIO24 informó pérdida de alimentación; también es compatible con poweroff manual local o señal espuria.",
            "¿Por qué no puede identificarse el origen exacto? Porque el journal anterior no persiste y la alerta no registra reason_code, boot ID, uptime ni motivo de apagado.",
        ],
        five_whys_id,
    )

    doc.add_heading("5. Impacto y respuesta de seguridad", level=1)
    add_bullets(
        doc,
        [
            "El fail-safe funcionó: el modo manual se abortó y el relé pasó a desenergizado en menos de 10 ms.",
            "Primera interrupción: aproximadamente 57 s hasta que el edge volvió a asignarse.",
            "Segunda interrupción: aproximadamente 1 h 53 min 50 s hasta que el edge volvió a asignarse.",
            "No se encontró evidencia de un despacho coincidente con los cortes.",
            "La clasificación de la alarma desvió el diagnóstico hacia el validador en vez de energía/apagado del edge.",
        ],
        bullet_id,
    )

    doc.add_heading("6. Acciones correctivas recomendadas", level=1)
    doc.add_heading("P0 - Terreno y alimentación", level=2)
    p0_id = add_numbering_definition(doc, ordered=True)
    add_numbered(
        doc,
        [
            "Contrastar 19:19 y 20:12 con maniobras del tablero, recarga del estanque, microcortes, protecciones y registros de la fuente 12-24 V.",
            "Revisar apriete y continuidad de bornes, fuente, tierra, UPS, conector interno y señales GPIO23/GPIO24.",
            "Si la alimentación nunca se retiró, capturar GPIO24 con registrador u osciloscopio; una caída espuria provocaría la misma secuencia.",
        ],
        p0_id,
    )

    doc.add_heading("P1 - Software", level=2)
    p1_software_id = add_numbering_definition(doc, ordered=True)
    add_numbered(
        doc,
        [
            "Agregar una marca explícita de cierre intencional en MqttValidatorTransport; close() debe limpiar RPC y estado sin invocar validator_link_lost.",
            "Conservar el fail-safe y la alerta para desconexiones inesperadas y códigos MQTT anormales.",
            "Crear un evento distinto y de baja criticidad para edge_service_stopping o uno específico para edge_power_lost.",
            "Enriquecer alertas con motivo, reason_code, boot ID, uptime, estado del sistema y origen del evento.",
        ],
        p1_software_id,
    )

    doc.add_heading("P1 - Observabilidad", level=2)
    p1_obs_id = add_numbering_definition(doc, ordered=True)
    add_numbered(
        doc,
        [
            "Cambiar journald a almacenamiento persistente con límite y rotación.",
            "Agregar un hook UPS que registre durablemente fecha, boot ID y GPIO24 antes del apagado.",
            "Publicar eventos diferenciados de pérdida y restauración de alimentación.",
        ],
        p1_obs_id,
    )

    doc.add_heading("P2 - Pruebas de aceptación", level=2)
    checks_id = add_numbering_definition(doc, ordered=True)
    add_numbered(
        doc,
        [
            "systemctl restart fuel-edge no crea una alerta de desconexión.",
            "systemctl poweroff registra cierre planificado sin crear una alerta del validador.",
            "Una caída MQTT inesperada sí corta el relé y crea la alerta correspondiente.",
            "Una prueba controlada de pérdida de alimentación registra el evento UPS y recupera el servicio al volver la energía.",
        ],
        checks_id,
    )

    doc.add_page_break()
    doc.add_heading("7. Estado al finalizar la revisión", level=1)
    add_status_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run("Nota: ").bold = True
    p.add_run(
        "vcgencmd get_throttled=0x0 corresponde al boot actual y no permite descartar condiciones eléctricas de boots anteriores."
    )

    doc.add_heading("8. Limitación forense", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "No se modificó configuración ni software productivo durante el RCA. El host usa journald con Storage=volatile, por lo que los logs detallados de los dos boots afectados se perdieron. "
    )
    p.add_run(
        "La relación causal entre apagado del edge y las dos alertas está demostrada; el origen eléctrico o humano exacto del flanco de apagado no puede establecerse al 100 %."
    ).bold = True

    doc.add_heading("9. Referencias", level=1)
    p = doc.add_paragraph(style="List Bullet")
    apply_numbering(p, bullet_id)
    add_hyperlink(
        p,
        "Industrial Shields - Raspberry Pi PLC: setup guide for UPS & RTC features",
        "https://www.industrialshields.com/blog/raspberry-pi-for-industry-26/how-to-work-with-ups-and-rtc-in-raspberry-plc-645",
    )
    p = doc.add_paragraph(style="List Bullet")
    apply_numbering(p, bullet_id)
    add_hyperlink(
        p,
        "Industrial Shields - GateBerry documentation",
        "https://docs.industrialshields.com/gateberry/",
    )
    p = doc.add_paragraph(style="List Bullet")
    apply_numbering(p, bullet_id)
    p.add_run("Base operacional sanitizada y auditoría local del edge, consultadas el 23-08-2026.")

    doc.add_heading("10. Convención del registro RCA", level=1)
    p = doc.add_paragraph()
    p.add_run("Formato: ").bold = True
    p.add_run("RCA-AAAA-NNN.")
    add_bullets(
        doc,
        [
            "AAAA identifica el año de emisión.",
            "NNN es un correlativo anual de tres dígitos, comenzando en 001.",
            f"Documento actual: {RCA_ID}.",
            "Próximo correlativo reservado: RCA-2026-002.",
            "Registro maestro: outputs/rca/registro-rca.csv.",
        ],
        bullet_id,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
